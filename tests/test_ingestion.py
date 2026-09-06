import io
import pytest
import fitz
import rag_ingestion as ingestion


def test_numeric_lines_and_unicode_survive_cleaning():
    text = 'Grade\n100\n2026\n🦕\nதமிழ்'
    assert ingestion._clean_text(text) == text


def test_utf16_and_bom_text():
    assert ingestion.convert_file_to_markdown('Notes: café'.encode('utf-16'), 'notes.txt')[0] == 'Notes: café'
    assert ingestion.convert_file_to_markdown('Notes'.encode('utf-8-sig'), 'notes.md')[0] == 'Notes'


@pytest.mark.parametrize('filename,data', [('app.exe', b'binary'), ('fake.txt', b'a\x00b'), ('blank.md', b''), ('photo.png', b'image'), ('broken.docx', b'not a zip')])
def test_bad_inputs_are_rejected(filename, data):
    with pytest.raises(ValueError):
        ingestion.convert_file_to_markdown(data, filename)


def test_csv_preserves_identifier_zeros():
    text, _ = ingestion.convert_file_to_markdown(b'Student,ID\nAsha,00123\n', 'students.csv')
    assert '00123' in text


def test_docx_preserves_table_order():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Before table')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Inside table'
    doc.add_paragraph('After table')
    buffer = io.BytesIO()
    doc.save(buffer)
    text, _ = ingestion.convert_file_to_markdown(buffer.getvalue(), 'notes.docx')
    assert text.index('Before table') < text.index('Inside table') < text.index('After table')


def test_xlsx_preserves_identifiers():
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['Name', 'ID'])
    book.active.append(['Asha', '00123'])
    buffer = io.BytesIO()
    book.save(buffer)
    text, _ = ingestion.convert_file_to_markdown(buffer.getvalue(), 'students.xlsx')
    assert '00123' in text


def test_pdf_password_is_reported():
    doc = fitz.open()
    doc.new_page().insert_text((50, 50), 'Private notes')
    data = doc.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw='owner', user_pw='password')
    with pytest.raises(ValueError, match='password'):
        ingestion.convert_file_to_markdown(data, 'private.pdf')


def test_mixed_pdf_ocr_only_reads_scanned_pages(monkeypatch):
    doc = fitz.open()
    doc.new_page().insert_text((50, 50), 'Readable first page')
    page = doc.new_page()
    page.draw_rect(fitz.Rect(10, 10, 100, 100), fill=(0.2, 0.2, 0.2))
    calls = []
    def fake_ocr(data, user_id, page_numbers):
        calls.append(page_numbers)
        return '## Page 2\n\nScanned second page'
    monkeypatch.setattr(ingestion, '_pdf_text_with_groq_vision', fake_ocr)
    text, _ = ingestion.convert_file_to_markdown(doc.tobytes(), 'mixed.pdf')
    assert 'Readable first page' in text and 'Scanned second page' in text
    assert calls == [[2]]


def test_blank_pdf_is_not_indexed():
    doc = fitz.open()
    doc.new_page()
    with pytest.raises(ValueError):
        ingestion.convert_file_to_markdown(doc.tobytes(), 'blank.pdf')
